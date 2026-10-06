from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "pacote-documentacao" / "Manual Completo Prime Acai PDV.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)

doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2.1)
section.right_margin = Cm(2.1)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.08
for name, size in [("Title", 25), ("Heading 1", 16), ("Heading 2", 12.5), ("Heading 3", 11)]:
    st = styles[name]
    st.font.name = "Aptos Display"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.font.bold = True
    st.paragraph_format.space_before = Pt(14 if name != "Title" else 0)
    st.paragraph_format.space_after = Pt(6)

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def margins(cell, top=100, start=100, bottom=100, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for idx, (cell, value) in enumerate(zip(t.rows[0].cells, headers)):
        cell.text = value
        shade(cell, "243A73")
        margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True
            run.font.size = Pt(9.5)
        if widths:
            cell.width = Cm(widths[idx])
    for r_idx, row in enumerate(rows):
        cells = t.add_row().cells
        for idx, (cell, value) in enumerate(zip(cells, row)):
            cell.text = str(value)
            margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if r_idx % 2:
                shade(cell, "F3F6FA")
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
            if widths:
                cell.width = Cm(widths[idx])
    doc.add_paragraph()
    return t

def bullets(items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")

def steps(items):
    for item in items:
        doc.add_paragraph(item, style="List Number")

title = doc.add_paragraph("Manual Completo do Sistema Prime Açaí PDV", style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
sub = doc.add_paragraph("Operação, administração, arquitetura, implantação e integrações")
sub.runs[0].font.size = Pt(13)
sub.runs[0].font.color.rgb = RGBColor(70, 78, 92)
doc.add_paragraph("Versão da documentação: 25 de setembro de 2026")
doc.add_paragraph(
    "Este manual consolida o funcionamento do cardápio digital e do painel de gestão da Açaí Prime. "
    "Ele orienta proprietários, operadores e responsáveis técnicos sobre pedidos, produtos, caixa, impressão, "
    "balança Prix 3 Toledo, preparação da integração PedeAI, banco de dados e publicação na Cloudflare."
)

doc.add_heading("1 Visão geral", level=1)
doc.add_paragraph(
    "O Prime Açaí PDV é um sistema web responsivo composto por cardápio público, painel administrativo, "
    "gestão de pedidos e mesas, banco Cloudflare D1 e um agente local para receber o peso da balança. "
    "A versão publicada mantém a identidade visual e o catálogo do cardápio original da Prime Açaí."
)
table(["Recurso", "Endereço ou identificação"], [
    ("Cardápio do cliente", "https://prime-acai.wpnz.com.br/"),
    ("Painel do proprietário", "https://prime-acai.wpnz.com.br/painel"),
    ("Página de acompanhamento", "https://prime-acai.wpnz.com.br/pedido"),
    ("Gestão de mesas", "https://prime-acai.wpnz.com.br/mesas"),
    ("Worker Cloudflare", "prime-acai-pdv"),
    ("Banco D1", "prime-acai-pdv-db"),
    ("Repositório", "wesleyxrodrigo19-sketch/prime-acai-pdv"),
], [5.2, 11.0])

doc.add_heading("2 Perfis de uso", level=1)
table(["Perfil", "Responsabilidades"], [
    ("Cliente", "Consulta o cardápio, seleciona complementos, informa entrega ou retirada e acompanha o pedido."),
    ("Atendente ou caixa", "Registra pedidos presenciais, recebe peso da balança, confere pagamentos e imprime comandas."),
    ("Produção", "Acompanha a fila, altera o status e sinaliza pedidos prontos."),
    ("Proprietário", "Administra produtos, disponibilidade, horários, taxas, mesas, relatórios e integrações."),
    ("Responsável técnico", "Mantém Cloudflare, D1, segredos, domínio, agente local e atualizações do código."),
], [4.0, 12.2])

doc.add_heading("3 Fluxos operacionais", level=1)
doc.add_heading("3.1 Pedido pelo cardápio", level=2)
steps([
    "O cliente abre o cardápio e escolhe produtos e complementos.",
    "O sistema calcula subtotal, taxa de entrega e total.",
    "O cliente informa identificação, telefone, modalidade, endereço e pagamento.",
    "O pedido é salvo no D1 e entra no painel com status Novo.",
    "A loja confirma, prepara, finaliza e conclui o atendimento.",
])
doc.add_heading("3.2 Pedido presencial por peso", level=2)
steps([
    "O atendente abre Novo pedido no balcão e localiza o cartão Self-service.",
    "O recipiente com o açaí é colocado na Prix 3 Toledo configurada em Prt1 ou Prt3.",
    "O agente local lê o peso pela porta USB serial e envia uma leitura estável ao PDV.",
    "O atendente seleciona Ler balança e adicionar. O sistema usa R$ 30,00 por kg, calcula o valor e adiciona o item à comanda.",
    "Antes de registrar, o servidor confere se o peso ainda corresponde à última leitura da balança e se foi capturado há no máximo dois minutos.",
    "A comanda registra peso em gramas, valor por kg e origem automática Prix 3 Toledo.",
])
doc.add_paragraph(
    "O preço inicial do self-service é R$ 30,00 por kg. O proprietário pode alterá-lo em "
    "Configurações, na seção Mesas e self-service, sem modificar o código do sistema."
)
doc.add_heading("3.3 Pedido de mesa", level=2)
doc.add_paragraph(
    "Pedidos presenciais podem ser vinculados a uma mesa. A quantidade de mesas é configurável no painel; "
    "o quadro mostra a situação e permite acompanhar consumo, pagamentos divididos e fechamento."
)

doc.add_heading("4 Cardápio do cliente", level=1)
bullets([
    "Busca por nome ou descrição e navegação por categorias.",
    "Produtos em destaque, imagens, preços, descrições e complementos.",
    "Bloqueio de novos pedidos quando a loja está fechada, mantendo consulta disponível.",
    "Entrega por bairro com taxa configurável, retirada e atendimento de mesa.",
    "Formas de pagamento: dinheiro, Pix, crédito e débito, conforme disponibilidade.",
    "Recuperação de carrinhos abandonados e registro de visitas para indicadores operacionais.",
])

doc.add_heading("5 Painel do proprietário", level=1)
doc.add_heading("5.1 Pedidos", level=2)
bullets([
    "Fila com status Novo, Confirmado, Preparando, Pronto, Saiu para entrega, Concluído e Cancelado.",
    "Identificação do canal, incluindo Delivery PedeAI quando a integração for ativada.",
    "Inclusão e alteração de itens, cancelamento, exclusão auditada, reimpressão e registro de pagamento.",
    "Detalhes de cliente, endereço, mesa, observações, troco, divisão de pagamento e totais.",
])
doc.add_heading("5.2 Produtos e categorias", level=2)
bullets([
    "Cadastro, edição, disponibilidade, imagem, preço, destaque, categoria e ordem de exibição.",
    "Complementos configuráveis por produto.",
    "Limite configurável de produtos destacados.",
    "O catálogo atual possui 29 produtos sincronizados com o cardápio original da Prime Açaí.",
])
doc.add_heading("5.3 Configurações", level=2)
bullets([
    "Dados da loja, horário comercial e fechamento manual com motivo.",
    "Taxas por bairro e acréscimo percentual opcional.",
    "Quantidade de mesas, preço do self-service por kg, cópias de impressão e parâmetros de destaques.",
    "Configuração da PedeAI e estado da última captura da balança.",
])
doc.add_heading("5.4 Relatórios e auditoria", level=2)
bullets([
    "Receita por período e formas de pagamento.",
    "Resumo de pedidos, visitas, conversão e carrinhos abandonados.",
    "Log de auditoria para ações sensíveis em pedidos.",
    "Estado dos trabalhos de impressão e falhas do agente local.",
])

doc.add_heading("6 Balança Prix 3 Toledo", level=1)
doc.add_paragraph(
    "O arquivo agents/prix3-bridge.ps1 executa no computador do caixa. Ele abre a porta serial, envia ENQ, "
    "interpreta a resposta STX mais peso mais ETX e transmite o peso em gramas ao endpoint protegido do Worker."
)
table(["Parâmetro", "Valor recomendado"], [
    ("Protocolo", "Prt1 ou Prt3"),
    ("Conexão", "USB serial por porta COM"),
    ("Baud rate", "9600, ajustável no script"),
    ("Formato serial", "8 bits, sem paridade, 1 stop bit"),
    ("Endpoint", "POST /api/scale/weight"),
    ("Autenticação", "Cabeçalho x-scale-key"),
    ("Preço inicial do self-service", "R$ 30,00 por kg, configurável no painel"),
    ("Validade da leitura", "Até 2 minutos; peso antigo ou divergente é recusado"),
], [5.0, 11.2])
doc.add_heading("6.1 Instalação do agente", level=2)
steps([
    "Conecte o cabo USB serial e confirme a porta COM no Gerenciador de Dispositivos.",
    "Configure a balança no protocolo Prt1 ou Prt3.",
    "Defina na Cloudflare o segredo SCALE_BRIDGE_KEY.",
    "No computador do caixa, execute o script informando a porta e a mesma chave.",
    "Coloque um peso conhecido, verifique a mensagem Peso transmitido e confira a leitura no painel.",
])
doc.add_paragraph("Exemplo: .\\prix3-bridge.ps1 -Porta COM3 -Chave SUA_CHAVE_SEGURA")
doc.add_heading("6.2 Diagnóstico", level=2)
table(["Sintoma", "Verificação"], [
    ("Porta não encontrada", "Confirme driver, cabo e número da COM."),
    ("Sem leitura", "Confira Prt1 ou Prt3, baud rate e estabilidade do peso."),
    ("Erro 401", "A chave local difere do segredo SCALE_BRIDGE_KEY."),
    ("Peso incorreto", "Revise a unidade retornada e o divisor configurado."),
    ("Painel não atualiza", "Verifique internet, endpoint e horário de captura."),
], [5.0, 11.2])

doc.add_heading("7 Integração PedeAI", level=1)
doc.add_paragraph(
    "O painel já contém a configuração necessária para receber credenciais, mas a integração operacional depende "
    "das credenciais e do contrato técnico fornecidos pela PedeAI. Enquanto estiver desativada, o webhook responde "
    "como integração em preparação."
)
table(["Campo", "Finalidade"], [
    ("Ambiente", "Sandbox ou Produção."),
    ("ID da loja", "Identificação da Açaí Prime na PedeAI."),
    ("Client ID", "Identificação da aplicação integradora."),
    ("URL base", "Endereço oficial da API PedeAI."),
    ("Token", "Autorização das chamadas à API."),
    ("Segredo do webhook", "Validação da assinatura dos eventos recebidos."),
], [5.0, 11.2])
doc.add_paragraph("Webhook reservado: https://prime-acai.wpnz.com.br/api/integrations/pedeai/webhook")
doc.add_heading("7.1 Sequência para ativação", level=2)
steps([
    "Solicitar credenciais de Sandbox, documentação de eventos e regras de autenticação.",
    "Preencher Configurações PedeAI no painel sem compartilhar os segredos em mensagens ou capturas.",
    "Cadastrar o webhook na PedeAI e validar assinatura, idempotência e resposta HTTP.",
    "Homologar entrega, retirada, pagamentos, descontos, complementos, cancelamento e atualização de status.",
    "Trocar para Produção somente após aprovação formal e executar um pedido real acompanhado.",
])

doc.add_heading("8 Arquitetura técnica", level=1)
table(["Camada", "Tecnologia e responsabilidade"], [
    ("Interface", "HTML, CSS e JavaScript responsivos servidos pelo Worker."),
    ("Aplicação", "Cloudflare Worker prime-acai-pdv."),
    ("Dados", "Cloudflare D1 SQLite, binding DB."),
    ("Domínio", "prime-acai.wpnz.com.br na Cloudflare."),
    ("Mídia", "Imagens remotas do cardápio e arquivos incorporados em /media."),
    ("Agente local", "PowerShell para Prix 3 Toledo e agente de impressão."),
    ("Entrega contínua", "GitHub main conectado ao build do Cloudflare Workers."),
], [4.0, 12.2])
doc.add_heading("8.1 Estrutura do repositório", level=2)
table(["Diretório ou arquivo", "Conteúdo"], [
    ("src", "Aplicação web, estilos, painel e template do Worker."),
    ("drizzle", "Migrações sequenciais do banco D1."),
    ("agents", "Agente local da balança Prix 3."),
    ("media", "Mídia incorporada na publicação."),
    ("scripts", "Build, pré-visualização, testes e geração de documentação."),
    ("dist/server/index.js", "Artefato gerado e publicado; não deve ser editado manualmente."),
    ("wrangler.jsonc", "Nome do Worker, entrada, compatibilidade e binding D1."),
], [5.0, 11.2])

doc.add_heading("9 Banco de dados", level=1)
table(["Tabela", "Finalidade"], [
    ("settings", "Configurações operacionais, horários, taxas e integrações."),
    ("products", "Produtos, categorias, preços, imagens, complementos e disponibilidade."),
    ("orders", "Pedidos, cliente, canal, entrega, pagamento, status e totais."),
    ("order_items", "Itens, preços, quantidades, complementos e observações."),
    ("order_payments", "Pagamentos divididos por pedido."),
    ("print_jobs", "Fila e estado das impressões."),
    ("abandoned_carts", "Carrinhos recuperáveis e histórico de contato."),
    ("site_visits", "Visitas agregadas para indicadores."),
    ("audit_log", "Rastro de ações administrativas sensíveis."),
], [4.6, 11.6])
doc.add_paragraph(
    "As migrações ficam em drizzle e devem ser aplicadas em ordem numérica. Nunca repita uma migração já aplicada, "
    "não exclua o banco e não altere tabelas diretamente sem registrar uma nova migração."
)

doc.add_heading("10 API principal", level=1)
table(["Método e rota", "Uso"], [
    ("GET /api/catalog", "Catálogo público e configurações visíveis."),
    ("POST /api/orders", "Criação de pedido pelo cliente."),
    ("GET /api/orders/:code", "Acompanhamento público do pedido."),
    ("POST /api/scale/weight", "Recebimento do peso pelo agente local."),
    ("POST /api/integrations/pedeai/webhook", "Recebimento reservado de eventos PedeAI."),
    ("POST /api/admin/login", "Autenticação do proprietário."),
    ("GET /api/admin/orders", "Fila administrativa de pedidos."),
    ("PATCH /api/admin/orders/:id", "Mudança de status do pedido."),
    ("GET e POST /api/admin/products", "Consulta e cadastro administrativo de produtos."),
    ("PATCH /api/admin/settings", "Atualização das configurações da loja."),
    ("GET /api/print/jobs", "Fila consumida pelo agente de impressão."),
], [7.0, 9.2])

doc.add_heading("11 Implantação e atualização", level=1)
doc.add_heading("11.1 Requisitos", level=2)
bullets([
    "Node.js 22 ou superior, Git e acesso à conta Cloudflare.",
    "Projeto Cloudflare Worker prime-acai-pdv e banco D1 prime-acai-pdv-db.",
    "Segredos ADMIN_PIN, SESSION_SECRET e SCALE_BRIDGE_KEY configurados somente na Cloudflare.",
])
doc.add_heading("11.2 Publicação", level=2)
steps([
    "Atualize o código e crie uma nova migração se o banco mudar.",
    "Execute npm run build e npm run test:dashboard.",
    "Aplique migrações remotas com npm run cf:d1:migrate após conferir o banco selecionado.",
    "Envie a branch main para o GitHub; o Cloudflare executará o build conectado.",
    "Confira Deployments, abra o cardápio, entre no painel e execute um pedido de teste.",
])
doc.add_heading("11.3 Verificação pós publicação", level=2)
bullets([
    "Todas as categorias, imagens e valores aparecem corretamente.",
    "Login do painel funciona e a sessão expira de forma segura.",
    "Pedido de teste percorre os status e gera impressão quando configurada.",
    "Taxas, horários, mesa, pagamento e total estão corretos.",
    "Agente da balança transmite peso e o painel mostra horário recente.",
])

doc.add_heading("12 Segurança e continuidade", level=1)
bullets([
    "Nunca registrar PIN, token, segredo ou chave de balança no GitHub, em documentação ou em mensagens abertas.",
    "Manter ADMIN_PIN, SESSION_SECRET e SCALE_BRIDGE_KEY como segredos do Worker.",
    "Usar HTTPS e validar autenticação antes de qualquer rota administrativa.",
    "Criar backup antes de migrações relevantes e usar Time Travel do D1 para recuperação emergencial.",
    "Revogar e substituir qualquer credencial compartilhada acidentalmente.",
    "Revisar audit_log quando houver cancelamento, exclusão ou divergência de pedido.",
])
doc.add_heading("12.1 Rotina recomendada", level=2)
table(["Periodicidade", "Atividade"], [
    ("Diária", "Conferir pedidos, falhas de impressão e última leitura da balança."),
    ("Semanal", "Revisar produtos indisponíveis, horários, taxas e carrinhos abandonados."),
    ("Mensal", "Exportar indicadores, revisar auditoria e validar acessos do proprietário."),
    ("Antes de atualização", "Registrar versão atual, conferir migrações e criar ponto de recuperação."),
    ("Após atualização", "Executar o checklist funcional e monitorar erros do Worker."),
], [4.4, 11.8])

doc.add_heading("13 Solução de problemas", level=1)
table(["Problema", "Ação inicial"], [
    ("Cardápio não abre", "Verificar deployment, domínio, DNS e status do Worker."),
    ("Imagem não aparece", "Conferir image_url, origem da imagem e cache do navegador."),
    ("Painel rejeita acesso", "Conferir ADMIN_PIN e SESSION_SECRET na Cloudflare."),
    ("Pedido não aparece", "Consultar resposta da API, tabela orders e logs do Worker."),
    ("Impressão parada", "Checar print_jobs, agente local, chave e impressora padrão."),
    ("Peso não chega", "Checar COM, protocolo, chave, internet e horário da leitura."),
    ("Migração falha", "Parar, identificar a última migração aplicada e não repetir alterações manualmente."),
    ("Domínio bloqueado na rede", "Testar em outra rede; filtros corporativos podem bloquear domínios recém-registrados."),
], [5.0, 11.2])

doc.add_heading("14 Checklist de entrega", level=1)
bullets([
    "Domínio e certificado HTTPS ativos.",
    "Cardápio com 29 produtos, imagens e valores conferidos.",
    "Painel, pedidos, mesas, pagamentos e relatórios testados.",
    "Segredos Cloudflare definidos e ausentes do repositório.",
    "Agente Prix 3 instalado e validado no computador do caixa.",
    "Agente de impressão configurado quando aplicável.",
    "PedeAI mantida desativada até receber credenciais e concluir homologação.",
    "Responsáveis da loja treinados para abertura, fechamento e tratamento de falhas.",
])

doc.add_heading("15 Controle de versão", level=1)
table(["Data", "Versão", "Resumo"], [
    ("25/09/2026", "1.2", "Self-service por peso a R$ 30,00/kg, validação da Prix 3 e preço configurável no painel."),
    ("25/09/2026", "1.1", "Sincronização das imagens originais de 27 produtos e imagens próprias para dois itens sem foto."),
    ("25/09/2026", "1.0", "Sistema Prime Açaí, catálogo completo, painel, balança e preparação PedeAI."),
], [3.2, 2.4, 10.6])

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run("Prime Açaí PDV  |  Manual completo  |  25 de setembro de 2026")
run.font.size = Pt(8)
run.font.color.rgb = RGBColor(100, 108, 120)

doc.save(OUT)
print(OUT)
