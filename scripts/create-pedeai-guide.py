from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Guia Integracao PedeAI Acai Prime.docx"
OUT.parent.mkdir(exist_ok=True)

doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2.0)
section.right_margin = Cm(2.0)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(7)
for name, size, color in [("Title", 25, "182130"), ("Heading 1", 16, "8D1C24"), ("Heading 2", 12, "243A73")]:
    style = styles[name]
    style.font.name = "Aptos Display"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = True
    style.paragraph_format.space_before = Pt(15 if name != "Title" else 0)
    style.paragraph_format.space_after = Pt(7)

def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)

def add_heading(text, level=1):
    return doc.add_paragraph(text, style=f"Heading {level}")

def add_bullets(items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")

def add_steps(items):
    for item in items:
        doc.add_paragraph(item, style="List Number")

title = doc.add_paragraph("Integração PedeAI Açaí Prime", style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
subtitle = doc.add_paragraph("Guia de configuração e homologação do recebimento de pedidos")
subtitle.runs[0].font.size = Pt(13)
subtitle.runs[0].font.color.rgb = RGBColor(83, 96, 116)

doc.add_paragraph(
    "O painel da Açaí Prime já possui a área Configuração PedeAI e uma URL de webhook reservada. "
    "Este guia mostra quais dados devem ser solicitados ao PedeAI, onde preenchê-los e como conduzir os testes antes de liberar pedidos reais. "
    "A integração não deve ser ativada em produção antes da homologação formal."
)

add_heading("O que já está preparado")
add_bullets([
    "Seção Configuração PedeAI dentro de Configurações no painel do proprietário.",
    "Campos para ambiente, ID da loja, Client ID, URL da API, token e segredo de webhook.",
    "URL reservada para recebimento: https://prime-acai.wpnz.com.br/api/integrations/pedeai/webhook.",
    "Indicador de que token e segredo já foram cadastrados, sem exibir os valores novamente.",
    "Identificação planejada no painel de pedidos: Delivery (PedeAI).",
])

add_heading("Dados que devem ser obtidos com o PedeAI")
table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Light Shading Accent 1"
headers = ["Informação", "Onde será usada", "Observação"]
for cell, value in zip(table.rows[0].cells, headers):
    cell.text = value
    shade(cell, "243A73")
    for run in cell.paragraphs[0].runs:
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.bold = True
for row in [
    ("Ambiente", "Campo Ambiente", "Começar em Sandbox ou homologação."),
    ("ID da loja ou estabelecimento", "ID da loja no PedeAI", "Identifica a Açaí Prime na plataforma."),
    ("Client ID ou App ID", "Client ID ou App ID", "Credencial pública da aplicação."),
    ("URL base da API", "URL base da API", "Usar exatamente a URL informada pelo PedeAI."),
    ("Token de acesso", "Token de acesso", "Dado sigiloso. Não enviar por WhatsApp ou imprimir."),
    ("Segredo ou assinatura de webhook", "Segredo do webhook", "Necessário quando a plataforma assina os eventos."),
    ("Documentação de eventos", "Equipe técnica", "Precisamos do formato dos pedidos, status e autenticação."),
]:
    cells = table.add_row().cells
    for cell, value in zip(cells, row):
        cell.text = value
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

add_heading("Passo a passo")
add_heading("1 Solicitar o acesso de integração", 2)
doc.add_paragraph(
    "Solicite ao PedeAI o cadastro da software house e as credenciais de Sandbox. Peça também a documentação técnica atualizada, o catálogo de eventos, os métodos de autenticação e os critérios de homologação."
)
add_heading("2 Preencher a configuração no painel", 2)
add_steps([
    "Abra https://prime-acai.wpnz.com.br/painel e entre no painel do proprietário.",
    "Acesse Configurações e localize a seção Configuração PedeAI.",
    "Selecione Sandbox ou homologação.",
    "Preencha ID da loja, Client ID, URL base, token e segredo, conforme os dados recebidos.",
    "Clique em Salvar configurações PedeAI. Os campos de token e segredo ficarão ocultos depois de salvos.",
    "Copie a URL de recebimento exibida e encaminhe-a ao suporte técnico do PedeAI quando for solicitado.",
])
add_heading("3 Homologar os fluxos", 2)
doc.add_paragraph("Execute os cenários abaixo em Sandbox. Registre data, número do pedido de teste, resultado e qualquer erro para enviar ao suporte.")
add_bullets([
    "Pedido de entrega com endereço, taxa e observação.",
    "Pedido para retirada no balcão.",
    "Pagamento em dinheiro com troco, cartão na entrega e pagamento online.",
    "Cupom, desconto de item, complemento ou adicional quando esses recursos estiverem habilitados.",
    "Alteração de status: novo, confirmado, preparando, pronto, saiu para entrega, concluído e cancelado.",
    "Recusa ou cancelamento de pedido e o retorno do status para o PedeAI.",
])
add_heading("4 Liberar produção", 2)
add_steps([
    "Receba do PedeAI a confirmação de homologação e as credenciais de Produção.",
    "Troque o Ambiente para Produção e atualize as credenciais no painel.",
    "Faça um pedido real acompanhado pelo responsável da loja.",
    "Confirme se ele aparece no quadro de pedidos como Delivery (PedeAI), com itens, valores, endereço e pagamento corretos.",
    "Ative o recebimento somente após essa conferência final.",
])

add_heading("Regras de operação e segurança")
add_bullets([
    "Nunca compartilhar token ou segredo de webhook em grupos, prints ou mensagens abertas.",
    "Manter o ambiente Sandbox enquanto a equipe PedeAI não aprovar a integração.",
    "Se uma credencial for trocada pelo PedeAI, substituir o valor no painel e salvar novamente.",
    "Antes de ativar Produção, confirmar que cardápio, disponibilidade, taxas de entrega e formas de pagamento no PedeAI coincidem com os dados da Açaí Prime.",
    "Definir quem acompanhará pedidos que falharem, forem cancelados ou apresentarem divergência de valor.",
])

add_heading("Ponto técnico para a homologação")
doc.add_paragraph(
    "A URL reservada responde que a integração está em preparação enquanto o recebimento estiver desligado. Após o PedeAI aprovar o formato técnico dos eventos, o conector será habilitado para transformar cada pedido recebido em uma comanda no painel, registrando a origem Delivery (PedeAI) e devolvendo as atualizações de status exigidas pela plataforma."
)

add_heading("Referência oficial")
doc.add_paragraph(
    "PedeAI. Introdução para integração: cadastro de software house, acesso Sandbox, homologação e ativação em produção. Disponível em https://pedeai.readme.io/reference/introdução. Acesso em 25 de setembro de 2026."
)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("Açaí Prime  |  Guia de integração PedeAI")
footer.runs[0].font.size = Pt(8)
footer.runs[0].font.color.rgb = RGBColor(108, 117, 129)

doc.save(OUT)
print(OUT)
